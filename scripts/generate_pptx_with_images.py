import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    IMG_DIR = os.path.join(BASE_DIR, "public", "presentation", "images")

    # Color Palette
    BG_COLOR = RGBColor(11, 15, 25)         # #0B0F19 Dark Navy
    CARD_BG = RGBColor(19, 27, 46)          # #131B2E Card Dark
    CARD_BORDER = RGBColor(40, 53, 80)      # Border
    TEXT_WHITE = RGBColor(248, 250, 252)    # Title White
    TEXT_MUTED = RGBColor(148, 163, 184)    # Slate Muted
    ACCENT_BLUE = RGBColor(59, 130, 246)    # Electric Blue
    ACCENT_CYAN = RGBColor(6, 182, 212)     # Cyan
    ACCENT_GREEN = RGBColor(16, 185, 129)   # Emerald Green
    ACCENT_AMBER = RGBColor(245, 158, 11)   # Amber Gold
    ACCENT_PURPLE = RGBColor(139, 92, 246)  # Purple

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_COLOR
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, eyebrow_text, title_text, subtitle_text=""):
        # Eyebrow
        eb_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        eb_tf = eb_box.text_frame
        eb_tf.word_wrap = True
        eb_tf.margin_top = eb_tf.margin_bottom = eb_tf.margin_left = eb_tf.margin_right = 0
        p_eb = eb_tf.paragraphs[0]
        p_eb.text = eyebrow_text.upper()
        p_eb.font.size = Pt(11)
        p_eb.font.bold = True
        p_eb.font.color.rgb = ACCENT_CYAN
        p_eb.font.name = "Calibri"

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.7))
        t_tf = t_box.text_frame
        t_tf.word_wrap = True
        t_tf.margin_top = t_tf.margin_bottom = t_tf.margin_left = t_tf.margin_right = 0
        p_t = t_tf.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.font.name = "Calibri"

        # Subtitle
        if subtitle_text:
            s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.38), Inches(11.7), Inches(0.55))
            s_tf = s_box.text_frame
            s_tf.word_wrap = True
            s_tf.margin_top = s_tf.margin_bottom = s_tf.margin_left = s_tf.margin_right = 0
            p_s = s_tf.paragraphs[0]
            p_s.text = subtitle_text
            p_s.font.size = Pt(12)
            p_s.font.color.rgb = TEXT_MUTED
            p_s.font.name = "Calibri"

    def add_card(slide, left, top, width, height, title, body_text="", tag="", tag_color=ACCENT_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        if tag:
            p_tag = tf.paragraphs[0]
            p_tag.text = tag.upper()
            p_tag.font.size = Pt(9.5)
            p_tag.font.bold = True
            p_tag.font.color.rgb = tag_color
            p_tag.font.name = "Calibri"
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Calibri"

        if body_text:
            p_body = tf.add_paragraph()
            p_body.text = body_text
            p_body.font.size = Pt(11)
            p_body.font.color.rgb = TEXT_MUTED
            p_body.font.name = "Calibri"

        return card

    def add_framed_picture(slide, img_name, left, top, width, height):
        img_path = os.path.join(IMG_DIR, img_name)
        if not os.path.exists(img_path):
            print(f"Warning: Image {img_name} not found at {img_path}")
            return None
        
        # Add background frame/border
        border_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left - Inches(0.04), top - Inches(0.04), width + Inches(0.08), height + Inches(0.08))
        border_box.fill.solid()
        border_box.fill.fore_color.rgb = RGBColor(30, 41, 65)
        border_box.line.color.rgb = RGBColor(59, 130, 246)
        border_box.line.width = Pt(1.5)

        pic = slide.shapes.add_picture(img_path, left, top, width=width, height=height)
        return pic

    # ==========================================
    # SLIDE 1: Title & Hero Banner (with Desktop Dashboard Image)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    
    # Left Hero Text
    h_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(5.6), Inches(3.2))
    h_tf = h_box.text_frame
    h_tf.word_wrap = True
    p1 = h_tf.paragraphs[0]
    p1.text = "AURA HRMS ENTERPRISE"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN

    p2 = h_tf.add_paragraph()
    p2.text = "Next-Gen Workforce Operating System"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    p3 = h_tf.add_paragraph()
    p3.text = "Unified Web Commander + Native iOS & Android Mobile Suite + Live GPS Travel Telemetry & Statutory Payroll."
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(10)

    # 2 Key Value Badges on Left
    add_card(slide1, Inches(0.8), Inches(4.7), Inches(2.7), Inches(2.1), "100% Automated Payroll", "Auto-calculates PF, PT, TDS & LOP deductions with single-click bulk payslips.", "Turnkey Compliance", ACCENT_GREEN)
    add_card(slide1, Inches(3.7), Inches(4.7), Inches(2.7), Inches(2.1), "Live Field Telemetry", "Real-time satellite GPS tracking of employee trips & linked expense receipts.", "Zero Fake Claims", ACCENT_BLUE)

    # Right: Hero Screenshot (03_hr_control_center.png)
    add_framed_picture(slide1, "03_hr_control_center.png", Inches(6.7), Inches(1.2), Inches(5.8), Inches(5.6))


    # ==========================================
    # SLIDE 2: Native Mobile Application Suite (Real Mobile Screenshots!)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Native Mobile Suite (iOS & Android)", "Employee Self-Service, Geofenced Attendance & Field Operations", "Equip your workforce with a fast, modern mobile app. Punch attendance, apply for leaves, download payslips, and track business trips on the go.")

    # 4 Real Mobile Phone Screenshots side-by-side
    mobile_imgs = [
        ("mobile_attendance.jpg", "1-Tap Geofenced Punch", Inches(0.8)),
        ("mobile_dashboard.jpg", "Mobile Self-Service Hub", Inches(3.8)),
        ("mobile_trips.jpg", "Field Trip & Mission Logs", Inches(6.8)),
        ("mobile_profile.jpg", "Employee ID & Leave Vault", Inches(9.8))
    ]

    for img_name, label, left_pos in mobile_imgs:
        add_framed_picture(slide2, img_name, left_pos, Inches(2.2), Inches(2.7), Inches(4.5))
        # Label below image
        lbl_box = slide2.shapes.add_textbox(left_pos, Inches(6.8), Inches(2.7), Inches(0.4))
        lbl_tf = lbl_box.text_frame
        lbl_tf.margin_top = lbl_tf.margin_bottom = lbl_tf.margin_left = lbl_tf.margin_right = 0
        p_lbl = lbl_tf.paragraphs[0]
        p_lbl.text = label
        p_lbl.font.size = Pt(11)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = ACCENT_CYAN
        p_lbl.alignment = PP_ALIGN.CENTER


    # ==========================================
    # SLIDE 3: Live Field Travel & GPS Telemetry (Real Route Screenshot!)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Market Differentiator", "Live Field Travel & Interactive GPS Route Telemetry", "Monitor sales reps, field engineers, and executive business trips in real time. Eliminate fake mileage and fraudulent meal vouchers.")

    # Left: Features
    add_card(slide3, Inches(0.8), Inches(2.2), Inches(4.5), Inches(1.5), "Live HR Satellite Operations Map", "View active traveling personnel with interactive GPS breadcrumb trails, vehicle routes, and destination milestones.", "Live Telemetry", ACCENT_BLUE)
    add_card(slide3, Inches(0.8), Inches(3.9), Inches(4.5), Inches(1.5), "Trip-Locked Expense Reimbursements", "Reimbursements are strictly attached to verified travel route IDs with camera receipt captures. No receipts, no travel movement = no reimbursement.", "Fraud Elimination", ACCENT_GREEN)
    add_card(slide3, Inches(0.8), Inches(5.6), Inches(4.5), Inches(1.4), "Executive Mission Dossiers", "Generates comprehensive travel dossiers with boarding passes, hotel bookings, and client deliverables.", "Complete Dossier", ACCENT_PURPLE)

    # Right: Real Telemetry Map Screenshot (09_trip_route_telemetry.png)
    add_framed_picture(slide3, "09_trip_route_telemetry.png", Inches(5.6), Inches(2.2), Inches(6.9), Inches(4.8))


    # ==========================================
    # SLIDE 4: Payroll Processing & Tax Compliance (Real Payroll Screenshot!)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Statutory Indian Payroll", "Single-Click Bulk Payroll Processing & Compliance Engine", "Automatically calculate PF, Professional Tax, TDS slabs, and dynamic Loss-of-Pay (LOP) deductions.")

    # Left: Features
    add_card(slide4, Inches(0.8), Inches(2.2), Inches(4.5), Inches(1.5), "Dynamic LOP Salary Deductions", "Unpaid leaves (Loss-of-Pay) automatically deduct daily wages from base salary without tedious manual spreadsheet math.", "Zero Human Error", ACCENT_GREEN)
    add_card(slide4, Inches(0.8), Inches(3.9), Inches(4.5), Inches(1.5), "Single-Click Bulk Generation", "Generate 50 to 5,000+ payslips in a single click. The system audits tax thresholds, generates PDFs, and distributes them securely.", "Fast-Track Payroll", ACCENT_BLUE)
    add_card(slide4, Inches(0.8), Inches(5.6), Inches(4.5), Inches(1.4), "Encrypted Payslip Vault", "Employees view historical breakdown (Basic, HRA, Travel Allowance, Net In-Hand) on web or mobile.", "Employee Self-Service", ACCENT_AMBER)

    # Right: Real Payroll Screenshot (05_payroll_management.png)
    add_framed_picture(slide4, "05_payroll_management.png", Inches(5.6), Inches(2.2), Inches(6.9), Inches(4.8))


    # ==========================================
    # SLIDE 5: Performance Appraisals & 360 Reviews (Real Appraisal Screenshot!)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Performance Management", "Appraisal Review Cycles & 360 Monthly Review Dossiers", "Objective, continuous performance management that aligns employee outcomes with corporate revenue goals.")

    # Left: Features
    add_card(slide5, Inches(0.8), Inches(2.2), Inches(4.5), Inches(1.5), "Two-Stage Review Cycles", "Self-appraisal ratings (1 to 5) followed by managerial score and qualitative feedback notes for fair, documented reviews.", "Fair Evaluations", ACCENT_CYAN)
    add_card(slide5, Inches(0.8), Inches(3.9), Inches(4.5), Inches(1.5), "Objectives & Key Results (OKRs)", "Employees set measurable quarterly goals with checkable key results and progress bars to monitor milestone completions.", "Outcome Driven", ACCENT_BLUE)
    add_card(slide5, Inches(0.8), Inches(5.6), Inches(4.5), Inches(1.4), "360 Monthly Review Dossier (PDF)", "A one-click printable executive sheet combining attendance rate, leaves, trips, OKRs, and manager notes for 1-on-1 meetings.", "1-on-1 Executive Sheet", ACCENT_GREEN)

    # Right: Real Appraisal Screenshot (10_performance_appraisal.png)
    add_framed_picture(slide5, "10_performance_appraisal.png", Inches(5.6), Inches(2.2), Inches(6.9), Inches(4.8))


    # ==========================================
    # SLIDE 6: Compensation Simulation & Increment Letters (Real Screenshot!)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Compensation Engineering", "Real-Time CTC Simulation & Branded Increment Letters", "Automates salary hike math, promotions, and formal legal documentation seamlessly.")

    # Left: Features
    add_card(slide6, Inches(0.8), Inches(2.2), Inches(4.5), Inches(1.5), "Real-Time CTC Simulator", "Test hike percentages (10%, 15%) or target CTCs. The engine automatically recalculates Basic (50%), HRA, PF, Gratuity, and Net In-Hand.", "Instant Math", ACCENT_AMBER)
    add_card(slide6, Inches(0.8), Inches(3.9), Inches(4.5), Inches(1.5), "1-Click Master Record Sync", "Once HR approves a proposed hike or promotion, the employee's master salary, designation, and tax deductions update immediately.", "Automated Master", ACCENT_GREEN)
    add_card(slide6, Inches(0.8), Inches(5.6), Inches(4.5), Inches(1.4), "Branded Increment Letters", "Generates formal, print-ready increment and promotion letters on corporate letterhead with full component breakdown tables.", "Formal Documentation", ACCENT_BLUE)

    # Right: Real Compensation Screenshot (16_compensation_revision.png)
    add_framed_picture(slide6, "16_compensation_revision.png", Inches(5.6), Inches(2.2), Inches(6.9), Inches(4.8))


    # ==========================================
    # SLIDE 7: AI Talent Acquisition & Resume Screener (Real Screenshot!)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Artificial Intelligence", "AI Candidate Resume Screener & Policy RAG Engine", "Harness vector search and machine intelligence directly inside your HR operations.")

    # Left: Features
    add_card(slide7, Inches(0.8), Inches(2.2), Inches(4.5), Inches(1.5), "Automated Batch Resume Parser", "Upload 50+ candidate PDF resumes. The system parses work history, skills, and scores candidates against the Job Description in seconds.", "Fast Hiring", ACCENT_PURPLE)
    add_card(slide7, Inches(0.8), Inches(3.9), Inches(4.5), Inches(1.5), "AI Policy RAG Chatbot", "Employees ask questions about maternity leave, notice period, or travel perks. The vector AI engine cites exact handbook clauses instantly.", "80% Less HR Queries", ACCENT_CYAN)
    add_card(slide7, Inches(0.8), Inches(5.6), Inches(4.5), Inches(1.4), "Candidate Interview Bot", "Send automated screening tokens (/interview/:token) for AI-driven preliminary interviews before manager rounds.", "Automated Screening", ACCENT_BLUE)

    # Right: Real Recruitment Screenshot (17_ai_talent_acquisition.png)
    add_framed_picture(slide7, "17_ai_talent_acquisition.png", Inches(5.6), Inches(2.2), Inches(6.9), Inches(4.8))


    # ==========================================
    # SLIDE 8: Executive Workforce Analytics (Real Analytics Screenshot!)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Executive Intelligence", "Workforce Telemetry, Department Headcounts & Attrition Trends", "Gives executive leadership complete visibility into organizational health, salary bands, and attendance compliance.")

    # Left: Features
    add_card(slide8, Inches(0.8), Inches(2.2), Inches(4.5), Inches(1.5), "Hiring vs Attrition Trends", "Track 12-month rolling joins vs exits to detect turnover hotspots before they impact revenue.", "Retention Analytics", ACCENT_CYAN)
    add_card(slide8, Inches(0.8), Inches(3.9), Inches(4.5), Inches(1.5), "Department & Salary Band Split", "Instant breakdown of payroll distribution, average employee tenure, and age/gender demographics.", "Cost Control", ACCENT_GREEN)
    add_card(slide8, Inches(0.8), Inches(5.6), Inches(4.5), Inches(1.4), "Real-Time Attendance Rate", "Daily attendance compliance rate across all company branches and field teams.", "Operational Health", ACCENT_AMBER)

    # Right: Real Analytics Screenshot (15_executive_analytics.png)
    add_framed_picture(slide8, "15_executive_analytics.png", Inches(5.6), Inches(2.2), Inches(6.9), Inches(4.8))


    # ==========================================
    # SLIDE 9: Security & DPDP Act 2023 Compliance
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Security & Compliance", "Bank-Grade Privacy: DPDP Act 2023 & Zero Data Sharing", "Your employee records, salary data, and identity proofs are protected by strict legal and technical controls.")

    add_card(slide9, Inches(0.8), Inches(2.3), Inches(2.7), Inches(4.5), "Dedicated Tenant Isolation", "Each company operates in an isolated, partitioned database. Cross-tenant data leakage is architecturally impossible.", "Architecture", ACCENT_BLUE)
    add_card(slide9, Inches(3.8), Inches(2.3), Inches(2.7), Inches(4.5), "AES-256 & TLS 1.3", "Military-grade encryption at rest for databases and documents, with TLS 1.3 encryption across all network transmissions.", "Cryptographic Guard", ACCENT_GREEN)
    add_card(slide9, Inches(6.8), Inches(2.3), Inches(2.7), Inches(4.5), "Strict Non-Sharing Pledge", "We never sell, rent, monetize, or share client data with third parties or data brokers. The client retains 100% data ownership.", "Non-Monetization", ACCENT_CYAN)
    add_card(slide9, Inches(9.8), Inches(2.3), Inches(2.7), Inches(4.5), "Enterprise SaaS Agreement", "Comprehensive Master Subscription Agreement with liability caps, mutual hold-harmless defense, and New Delhi arbitration.", "Legal Shield", ACCENT_AMBER)


    # ==========================================
    # SLIDE 10: Competitor Benchmark Matrix
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Market Comparison", "How AURA HRMS Compares to Legacy Competitors", "High enterprise functionality without exorbitant multi-year contract traps.")

    table_shape = slide10.shapes.add_table(6, 5, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.6))
    table = table_shape.table

    cols = ["Core Feature", "AURA HRMS (Our Platform)", "Keka HR", "Darwinbox", "greytHR"]
    for i, name in enumerate(cols):
        cell = table.cell(0, i)
        cell.text = name
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(26, 36, 60)
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = ACCENT_CYAN if i == 1 else TEXT_WHITE

    rows_data = [
        ["Mobile App (iOS & Android)", "Included Free (Full Feature)", "Yes (Base features)", "Yes (Complex)", "Yes (Basic)"],
        ["Live Field GPS Map Telemetry", "Native Built-in Map", "Paid Add-on (High cost)", "Third-party plugin only", "Not Available"],
        ["AI Policy RAG & Resume Screener", "Built-in Core Engine", "Limited Chatbot", "Enterprise Custom Tier", "Not Available"],
        ["Implementation Timeline", "Under 24 Hours", "3 to 6 Weeks", "2 to 4 Months", "2 to 3 Weeks"],
        ["Pricing & Contract Lock-In", "Zero Lock-In (Monthly Billing)", "Min. Rs 60,000+ Annual Lock", "Min. Rs 2,50,000+ Annual Lock", "Min. Rs 35,000+ Annual Lock"]
    ]

    for row_idx, r_data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(r_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(16, 23, 40) if col_idx != 1 else RGBColor(24, 38, 70)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = ACCENT_GREEN if (col_idx == 1 and row_idx in [1, 2, 3, 4, 5]) else TEXT_MUTED


    # ==========================================
    # SLIDE 11: Risk-Free 30-Day Free Pilot & CTA
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    
    cta = add_card(slide11, Inches(1.5), Inches(1.2), Inches(10.3), Inches(5.2), "Experience Aura HRMS with a Risk-Free 30-Day Free Pilot", "", "Special Introductory Offer", ACCENT_GREEN)
    tf_cta = cta.text_frame
    
    p_desc = tf_cta.add_paragraph()
    p_desc.text = "Zero upfront setup fees. We configure your company's dedicated workspace, migrate your employee records, and deploy the native mobile app for your entire workforce completely free for 30 days."
    p_desc.font.size = Pt(15)
    p_desc.font.color.rgb = TEXT_MUTED
    p_desc.space_before = Pt(14)
    p_desc.space_after = Pt(20)

    bullets = [
        "Dedicated Account Onboarding Specialist assigned to your team",
        "Full access to Native Mobile Apps (iOS & Android) with live geofenced punch",
        "Complete Indian statutory payroll processing (PF, PT, TDS, LOP)",
        "Zero long-term contract lock-in — continue only if you love the platform"
    ]
    for b in bullets:
        p_b = tf_cta.add_paragraph()
        p_b.text = "   " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.font.bold = True
        p_b.space_after = Pt(8)

    # Save Deck
    output_path = os.path.join(BASE_DIR, "Aura_HRMS_Enterprise_Pitch_Deck.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
